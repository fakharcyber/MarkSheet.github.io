from openpyxl import Workbook, load_workbook
import os

file_Name = "student_up.xlsx"

if os.path.exists(file_Name):
    workbook=load_workbook(file_Name)
    sheet=workbook.active
else:
    workbook=Workbook()
    sheet=workbook.active
    sheet.title="students"
    sheet.append(["Roll_No","Name","Class","English","Urdu","Math","Chemistry","Physics","Total_Marks","Obt_Marks","Per%","Grade","Batch"])

no_of_stds= int(input("Enter the quentity of students: "))   


for i in range(no_of_stds):
   print("\nstudents",i+1)

   roll_No= input("Enter Roll_No: ")
   name= input("Enter Name: ")
   classs= input("Enter class: ")

   english = int(input("Enter English No: "))
   urdu = int(input("Enter Urdu No: "))
   math = int(input("Enter Math No: "))
   chemistry = int(input("Enter Chemistry No: "))
   physics = int(input("Enter Physics No: "))


   total_marks= 500
   obt_marks= english+urdu+math+chemistry+physics
   per= (obt_marks / total_marks)*100
   
   if min(english, urdu, math, chemistry, physics) < 35:
    Grade = "F"
   elif per >= 80:
       Grade = "A+"
   elif per >= 70:
       Grade = "A"
   elif per >= 60:
       Grade = "B"
   elif per >= 50:
       Grade = "C"
   elif per >= 40:
       Grade = "D"
   else: 
       Grade = "F"

   # Ye loop ke andar hona chahiye
   sheet.append([
       roll_No, name, classs, english, urdu,
       math, chemistry, physics, total_marks,
       obt_marks, round(per, 2), Grade, "Pr202312e"
   ])

# Ye loop ke bahar hona chahiye
workbook.save(file_Name)
print("===== Update Data =====")
 
     
   























